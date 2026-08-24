"""Excel 预览:.xlsx 依赖 openpyxl;.xls(OLE) 缺库则 unsupported,不当 zip。"""

from __future__ import annotations

import importlib.util
from pathlib import Path

import pytest

from insar_agent.preview.dispatch import PreviewOpts, preview_file
from insar_agent.preview.excel import preview_excel

HAS_OPENPYXL = importlib.util.find_spec("openpyxl") is not None
HAS_XLRD = importlib.util.find_spec("xlrd") is not None

_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


@pytest.mark.skipif(HAS_OPENPYXL, reason="openpyxl installed")
def test_xlsx_missing_openpyxl_unsupported(tmp_path: Path) -> None:
    path = tmp_path / "a.xlsx"
    path.write_bytes(b"not-a-workbook")
    result = preview_excel(path)
    assert result.kind == "unsupported"
    assert result.payload["reason"] == "missing_openpyxl"
    assert "openpyxl" in (result.note or "")


@pytest.mark.skipif(not HAS_OPENPYXL, reason="openpyxl not installed")
def test_xlsx_first_sheet_and_sheets_list(tmp_path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "First"
    ws1.append(["name", "value"])
    ws1.append(["alpha", "1"])
    ws2 = wb.create_sheet("Second")
    ws2.append(["other"])
    ws2.append(["zzz"])
    path = tmp_path / "tiny.xlsx"
    wb.save(path)
    wb.close()

    result = preview_excel(path)
    assert result.kind == "table"
    assert result.payload["sheet"] == "First"
    assert result.payload["sheets"] == ["First", "Second"]
    assert result.payload["columns"] == ["name", "value"]
    assert result.payload["rows"] == [["alpha", "1"]]
    assert result.payload["n_rows_total"] == 1
    assert result.payload["n_cols_total"] == 2
    assert result.truncated is False


@pytest.mark.skipif(not HAS_OPENPYXL, reason="openpyxl not installed")
def test_xlsx_switch_sheet(tmp_path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "First"
    ws1.append(["name"])
    ws1.append(["alpha"])
    ws2 = wb.create_sheet("Second")
    ws2.append(["other"])
    ws2.append(["zzz"])
    path = tmp_path / "tiny.xlsx"
    wb.save(path)
    wb.close()
    result = preview_file(path, PreviewOpts(sheet="Second"))
    assert result.kind == "table"
    assert result.payload["sheet"] == "Second"
    assert result.payload["columns"] == ["other"]
    assert result.payload["rows"] == [["zzz"]]


@pytest.mark.skipif(HAS_XLRD, reason="xlrd installed")
def test_xls_without_support_unsupported(tmp_path: Path) -> None:
    ole = tmp_path / "legacy.xls"
    ole.write_bytes(_OLE_MAGIC + b"\x00" * 32)
    zipped = tmp_path / "zipped.xls"
    zipped.write_bytes(b"PK\x03\x04" + b"\x00" * 32)

    for path in (ole, zipped):
        result = preview_excel(path)
        assert result.kind == "unsupported"
        assert result.payload.get("reason") == "missing_xlrd"
        assert result.note
        assert "zip" in result.note.lower() or "OLE" in (result.note or "")
