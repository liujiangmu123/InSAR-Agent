"""Excel 预览:.xlsx/.xlsm 走 openpyxl;.xls(OLE) 只走 xlrd,不当 zip 解。"""

from __future__ import annotations

from collections.abc import Iterable
from pathlib import Path
from typing import Any

from insar_agent.preview.dispatch import Preview, current_opts, register


def _cell_str(value: object) -> str:
    if value is None:
        return ""
    return str(value)


def _unsupported(note: str, reason: str) -> Preview:
    return Preview(kind="unsupported", note=note, payload={"reason": reason})


def _from_rows(
    row_iter: Iterable[Iterable[Any]],
    *,
    sheets: list[str],
    sheet: str,
) -> Preview:
    opts = current_opts()
    offset = opts.offset
    col_offset = opts.col_offset
    limit = opts.limit
    col_limit = opts.col_limit
    columns: list[str] | None = None
    rows: list[list[str]] = []
    n_rows_total = 0
    n_cols_total = 0
    for i, raw in enumerate(row_iter):
        row = list(raw)
        n_cols_total = max(n_cols_total, len(row))
        shown = [_cell_str(c) for c in row[col_offset: col_offset + col_limit]]
        if i == 0:
            columns = shown
            continue
        n_rows_total += 1
        if n_rows_total <= offset:
            continue
        if len(rows) < limit:
            rows.append(shown)
    if columns is None:
        columns = []
    page_end = offset + len(rows)
    truncated = (
        n_rows_total > page_end
        or offset > 0
        or n_cols_total > col_offset + len(columns)
        or col_offset > 0
    )
    if n_rows_total == 0 and not columns:
        note = "空工作表"
        truncated = False
    elif truncated:
        note = (
            f"工作表 {sheet} · 第 {offset + 1}–{max(offset, page_end)} 行"
            f"（共 {n_rows_total} 行 × {n_cols_total} 列,可翻页/切表）"
        )
    else:
        note = f"工作表 {sheet} · {n_rows_total} 行 × {n_cols_total} 列（本表本页完整）"
    return Preview(
        kind="table",
        payload={
            "columns": columns,
            "rows": rows,
            "n_rows_total": n_rows_total,
            "n_cols_total": n_cols_total,
            "sheets": sheets,
            "sheet": sheet,
            "offset": offset,
            "col_offset": col_offset,
            "limit": limit,
            "col_limit": col_limit,
        },
        truncated=truncated,
        note=note,
    )


def _preview_xlsx(path: Path) -> Preview:
    try:
        import openpyxl
    except ImportError:
        return _unsupported("未安装 openpyxl，无法预览 Excel", "missing_openpyxl")

    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — 坏文件不得伪装成表格
        return _unsupported(f"无法打开 Excel:{type(exc).__name__}", "invalid_excel")

    try:
        sheets = [str(name) for name in wb.sheetnames]
        if not sheets:
            ws = wb.active
            if ws is None:
                return _unsupported("工作簿没有工作表", "empty_workbook")
            sheets = [str(ws.title)]
        wanted = current_opts().sheet
        sheet = wanted if wanted in sheets else sheets[0]
        ws = wb[sheet]
        return _from_rows(ws.iter_rows(values_only=True), sheets=sheets, sheet=sheet)
    finally:
        wb.close()


def _preview_xls(path: Path) -> Preview:
    """OLE .xls:只用 xlrd。即使文件其实是 zip/xlsx,也不按 zip 解析。"""
    try:
        import xlrd
    except ImportError:
        return _unsupported(
            "未安装 xlrd，无法预览 .xls（OLE 格式，不会按 zip 解析）",
            "missing_xlrd",
        )

    try:
        book = xlrd.open_workbook(str(path), on_demand=True)
    except Exception as exc:  # noqa: BLE001 — 坏 OLE 不得伪装成表格
        return _unsupported(f"无法解析 .xls（OLE）:{type(exc).__name__}", "invalid_xls")

    try:
        sheets = [str(name) for name in book.sheet_names()]
        if not sheets:
            return _unsupported("工作簿没有工作表", "empty_workbook")
        wanted = current_opts().sheet
        if wanted in sheets:
            sh = book.sheet_by_name(wanted)
            sheet = wanted
        else:
            sh = book.sheet_by_index(0)
            sheet = sheets[0]

        def rows() -> Iterable[list[Any]]:
            for r in range(sh.nrows):
                yield [sh.cell_value(r, c) for c in range(sh.ncols)]

        return _from_rows(rows(), sheets=sheets, sheet=sheet)
    finally:
        book.release_resources()


@register(".xlsx", ".xlsm", ".xls")
def preview_excel(path: Path) -> Preview:
    if path.suffix.lower() == ".xls":
        return _preview_xls(path)
    return _preview_xlsx(path)
